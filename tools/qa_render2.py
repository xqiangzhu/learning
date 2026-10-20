#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render pptx slides (geometry + text) to PNG for visual QA."""
from pptx import Presentation
from pptx.util import Pt
from PIL import Image, ImageDraw, ImageFont

SCALE = 130  # px per inch
FONT_PATH = "/tmp/fonts-wqy-microhei-master/wqy-microhei.ttc"
SRC = "/home/user/learning/体检平台架构决策与核心流程.pptx"

_fcache = {}
def F(pt):
    key = int(pt * 4)
    if key not in _fcache:
        _fcache[key] = ImageFont.truetype(FONT_PATH, max(int(pt * SCALE / 72), 6))
    return _fcache[key]

def color_of(cf):
    try:
        r = cf.rgb
        if r is not None and str(r) != "None":
            return (r[0], r[1], r[2])
    except Exception:
        pass
    return None

def solid_fill(sh):
    try:
        if int(sh.fill.type) == 1:
            return color_of(sh.fill.fore_color)
    except Exception:
        pass
    return None

prs = Presentation(SRC)
W = int(prs.slide_width / 914400 * SCALE)
H = int(prs.slide_height / 914400 * SCALE)

def IN(emu):
    return emu / 914400 * SCALE

def draw_text_frame(d, sh, x, y, w, h):
    tf = sh.text_frame
    paras = []
    for p in tf.paragraphs:
        runs = [(r.text, (r.font.size.pt if r.font.size else 10),
                 bool(r.font.bold), color_of(r.font.color) or (36, 50, 78)) for r in p.runs]
        if runs:
            paras.append(runs)
    if not paras:
        return
    # measure
    line_h = []
    wrapped = []
    for runs in paras:
        size = max(r[1] for r in runs)
        f = F(size)
        text = "".join(r[0] for r in runs)
        maxw = w - 4
        lines = []
        cur = ""
        for ch in text:
            if f.getlength(cur + ch) > maxw and cur:
                lines.append(cur); cur = ch
            else:
                cur += ch
        lines.append(cur)
        for ln in lines:
            line_h.append(size * SCALE / 72 * 1.12)
        wrapped.append((lines, runs, size))
    total = sum(line_h)
    ty = y + max((h - total) / 2, 0)
    for lines, runs, size in wrapped:
        f = F(size)
        color = runs[0][3]
        bold = runs[0][2]
        for ln in lines:
            # crude align: center
            tw = f.getlength(ln)
            tx = x + max((w - tw) / 2, 2)
            d.text((tx, ty), ln, font=f, fill=color,
                   stroke_width=1 if bold else 0, stroke_fill=color)
            ty += size * SCALE / 72 * 1.12

for idx, slide in enumerate(prs.slides):
    img = Image.new("RGB", (W, H), (255, 255, 255))
    d = ImageDraw.Draw(img)
    # backgrounds first, then shapes, then connectors, then text
    texts = []
    for sh in slide.shapes:
        if sh.element.tag.endswith('}cxnSp'):
            continue
        x, y = IN(sh.left), IN(sh.top)
        w, h = IN(sh.width), IN(sh.height)
        fill = solid_fill(sh)
        line_c = color_of(sh.line.color) or (154, 170, 196)
        try:
            if int(sh.line.fill.type) == 5:  # background -> no line
                line_c = None
        except Exception:
            pass
        st = str(sh.shape_type)
        if "OVAL" in st:
            d.ellipse([x, y, x + w, y + h], fill=fill, outline=line_c, width=1)
        else:
            d.rectangle([x, y, x + w, y + h], fill=fill, outline=line_c, width=1)
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append((sh, x, y, w, h))
    for sh in slide.shapes:
        if sh.element.tag.endswith('}cxnSp'):
            x1, y1 = IN(sh.left), IN(sh.top)
            w, h = IN(sh.width), IN(sh.height)
            from pptx.oxml.ns import qn
            xfrm = sh.element.find('.//' + qn('a:xfrm'))
            flipH = xfrm is not None and xfrm.get('flipH') == '1'
            flipV = xfrm is not None and xfrm.get('flipV') == '1'
            xa, xb = (x1 + w, x1) if flipH else (x1, x1 + w)
            ya, yb = (y1 + h, y1) if flipV else (y1, y1 + h)
            d.line([xa, ya, xb, yb], fill=(81, 101, 133), width=2)
            d.polygon([(xb, yb), (xb - 6, yb - 3), (xb - 3, yb - 6)], fill=(200, 60, 50))
    for sh, x, y, w, h in texts:
        draw_text_frame(d, sh, x, y, w, h)
    img.save(f"/tmp/qa2_{idx+1:02d}.png")

# contact sheets 2x2
from math import ceil
n = len(prs.slides._sldIdLst)
for g in range(ceil(n / 4)):
    sheet = Image.new("RGB", (W // 2 * 2 + 12, (H // 2) * 2 + 12), (230, 232, 238))
    for k in range(4):
        i = g * 4 + k
        if i >= n: break
        im = Image.open(f"/tmp/qa2_{i+1:02d}.png").resize((W // 2, H // 2))
        sheet.paste(im, ((k % 2) * (W // 2 + 4), (k // 2) * (H // 2 + 4)))
    sheet.save(f"/tmp/qa2_grid_{g+1}.png")
print("done", n)
