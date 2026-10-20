# -*- coding: utf-8 -*-
"""
PPTX 渲染器：把 deck_content.SLIDES 渲染成 16:9 的 .pptx

设计约束
  * 不依赖任何模板文件，全部用 python-pptx 原生形状绘制，保证在任何机器上可编辑；
  * 中英文字体统一设为 Microsoft YaHei（同时写入 a:latin / a:ea，避免中文回退成宋体）；
  * 所有文本框都过一遍 em 宽度估算（FIT_LOG），构建后可用 verify 报告检查是否溢出；
  * 只使用 Microsoft YaHei 覆盖良好的符号：· ● ▪ □ ■ × → ← ↔ ①②③，不使用 ⇄ ⚡ 等。
"""
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn

from deck_content import (SLIDES, C, FONT, META, FOOTER, DW, DH, DOMAINS,
                          anchor_points)

# ------------------------------------------------------------------ 版面 ----
SW, SH = 13.333, 7.5                 # 16:9
ML = MR = 0.72
CW = SW - ML - MR                    # 11.893 内容宽
BODY_TOP_WITH_LEDE = 1.98
BODY_TOP_PLAIN = 1.62
BODY_BOT = 6.62
FN_Y = 6.72                          # 页脚注（footnote）
FT_Y = 7.10                          # 页脚（footer）

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "mid": MSO_ANCHOR.MIDDLE, "bot": MSO_ANCHOR.BOTTOM}

FIT_LOG = []                         # 溢出检查记录


# ------------------------------------------------------- 文本宽度估算工具 ----
def em_width(s):
    """按 em 估算字符串宽度（1.0 = 一个全角字符宽）。"""
    w = 0.0
    for ch in s:
        o = ord(ch)
        if (0x2E80 <= o <= 0xA4CF or 0xAC00 <= o <= 0xD7A3 or 0xF900 <= o <= 0xFAFF
                or 0xFE30 <= o <= 0xFE6F or 0xFF00 <= o <= 0xFF60 or 0xFFE0 <= o <= 0xFFE6
                or 0x2460 <= o <= 0x25FF or 0x3000 <= o <= 0x303F or 0x1100 <= o <= 0x115F):
            w += 1.0
        elif ch in "▪●□■":
            w += 1.0
        elif ch in "×→←↔·":
            w += 0.72
        elif ch.isupper():
            w += 0.70
        elif ch.isdigit():
            w += 0.56
        elif ch == " ":
            w += 0.30
        elif ch in ".,:;!'|()/_-":
            w += 0.34
        else:
            w += 0.52
    return w


def _expand_newlines(paras):
    """把 run 文本里的 \n 展开成独立段落，便于测量与渲染。"""
    out = []
    for p in paras:
        runs = p.get("runs")
        if runs is None:
            runs = [{"t": p.get("text", "")}]
        chunks = [[]]
        for r in runs:
            parts = str(r.get("t", "")).split("\n")
            for i, seg in enumerate(parts):
                if i:
                    chunks.append([])
                if seg:
                    rr = dict(r)
                    rr["t"] = seg
                    chunks[-1].append(rr)
        for ci, ch in enumerate(chunks):
            np = dict(p)
            np["runs"] = ch
            if ci:
                np["space_before"] = np.get("space_before", 0) or 2
            out.append(np)
    return out


def measure(paras, box_w_in, pad_in=0.0):
    """估算段落组渲染后的高度（英寸）。行高按 size * 1.22 * line_spacing。"""
    avail = max((box_w_in - pad_in) * 72.0, 1.0)
    h_pt = 0.0
    for p in paras:
        size = p.get("size", 12)
        ls = p.get("line_spacing", 1.16)
        wpt = sum(em_width(r.get("t", "")) * r.get("size", size) for r in p["runs"])
        lines = max(1, math.ceil(wpt / avail - 1e-6))
        h_pt += lines * size * 1.22 * ls
        h_pt += (p.get("space_before") or 0) + (p.get("space_after") or 0)
    return h_pt / 72.0


# ------------------------------------------------------------- 基础绘制 ----
def _set_ea(run, name):
    """同时设置西文与东亚字体，避免中文被回退成宋体。"""
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def _set_spc(run, spc):
    """字符间距（1/100 pt）。"""
    run._r.get_or_add_rPr().set("spc", str(int(spc)))


def _no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def rgb(key_or_hex):
    return RGBColor.from_string(C.get(key_or_hex, key_or_hex))


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, rounded=True,
         radius=0.10, dash=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(max(w, 0.01)), Inches(max(h, 0.01)))
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(lw)
        if dash:
            shp.line.dash_style = dash
    _no_shadow(shp)
    shp.text_frame.word_wrap = True
    return shp


def line(slide, x1, y1, x2, y2, color="line", w=0.75, dash=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(w)
    if dash:
        conn.line.dash_style = dash
    _no_shadow(conn)
    return conn


def arrow(slide, x1, y1, x2, y2, color="muted", w=1.2, dash=None):
    conn = line(slide, x1, y1, x2, y2, color, w, dash)
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return conn


def text(slide, x, y, w, h, paras, anchor="top", wrap=True, page=0, where="",
         pad=0.0):
    """paras: [{text|runs, size, bold, color, align, line_spacing, space_before/after}]"""
    paras = _expand_newlines(paras)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(max(w, 0.05)),
                                  Inches(max(h, 0.05)))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]
    first = True
    for p in paras:
        par = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        par.alignment = ALIGN[p.get("align", "left")]
        par.line_spacing = p.get("line_spacing", 1.16)
        if p.get("space_before"):
            par.space_before = Pt(p["space_before"])
        if p.get("space_after") is not None:
            par.space_after = Pt(p["space_after"])
        size = p.get("size", 12)
        for r in p["runs"]:
            run = par.add_run()
            run.text = r.get("t", "")
            f = run.font
            fname = r.get("font", p.get("font", FONT))
            f.name = fname
            f.size = Pt(r.get("size", size))
            f.bold = r.get("bold", p.get("bold", False))
            f.italic = r.get("italic", p.get("italic", False))
            f.color.rgb = rgb(r.get("color", p.get("color", "ink")))
            _set_ea(run, fname)
            if r.get("spc") or p.get("spc"):
                _set_spc(run, r.get("spc") or p.get("spc"))
    need = measure(paras, w, pad)
    FIT_LOG.append({"page": page, "where": where, "w": round(w, 2), "h": round(h, 2),
                    "need": round(need, 2),
                    "over": round(need - h, 2),
                    "head": "".join(r.get("t", "") for r in paras[0]["runs"])[:34]})
    return tb


def P(t="", **kw):
    """快捷构造段落。"""
    p = {"runs": [{"t": t}]}
    p.update(kw)
    return p


def R(t, **kw):
    r = {"t": t}
    r.update(kw)
    return r


def pill(slide, x, y, w, h, label, fill, color="white", size=9, bold=True,
         line_color=None, align="center"):
    rect(slide, x, y, w, h, fill=fill, line=line_color, lw=0.75, rounded=True,
         radius=0.5)
    text(slide, x, y + (h - size * 1.22 / 72) / 2 - 0.012, w, size * 1.3 / 72,
         [P(label, size=size, bold=bold, color=color, align=align)],
         anchor="mid")


# ------------------------------------------------------------ 页眉 / 页脚 ----
def header(slide, spec, page, dark=False):
    kicker = spec.get("kicker")
    title = spec.get("title")
    lede = spec.get("lede")
    kc = "accent" if not dark else "accent"
    tc = "white" if dark else "ink"
    y = 0.44
    if kicker:
        text(slide, ML, y, CW, 0.24,
             [P(kicker, size=10, bold=True, color=kc, spc=110)], page=page,
             where="kicker")
        y = 0.70
    if title:
        text(slide, ML, y, CW, 0.52,
             [P(title, size=25, bold=True, color=tc, line_spacing=1.0)],
             page=page, where="title")
    rect(slide, ML, 1.28, 0.66, 0.045, fill="accent", rounded=False)
    line(slide, ML + 0.78, 1.303, ML + CW, 1.303, color="navy2" if dark else "line",
         w=0.75)
    if lede:
        text(slide, ML, 1.40, CW - 0.2, 0.5,
             [P(lede, size=11.5, color="CBD6E2" if dark else "muted",
                line_spacing=1.2)], page=page, where="lede")
        return BODY_TOP_WITH_LEDE
    return BODY_TOP_PLAIN


def footer(slide, page, total, dark=False):
    c = "8FA3B8" if dark else "muted"
    line(slide, ML, FT_Y - 0.06, ML + CW, FT_Y - 0.06,
         color="navy2" if dark else "line", w=0.75)
    text(slide, ML, FT_Y, CW - 1.2, 0.24, [P(FOOTER, size=8.5, color=c)])
    text(slide, ML + CW - 1.2, FT_Y, 1.2, 0.24,
         [P("%02d / %d" % (page, total), size=8.5, color=c, align="right")])


def bg(slide, color="bg"):
    rect(slide, 0, 0, SW, SH, fill=color, rounded=False)


def body_geom(spec):
    """返回 (top, bottom, callout_rect|None) —— callout 占位后主内容区自动收缩。"""
    top = BODY_TOP_WITH_LEDE if spec.get("lede") else BODY_TOP_PLAIN
    bot = BODY_BOT
    cal = spec.get("callout")
    if cal:
        ch = 0.46 if len(cal["text"]) < 62 else 0.62
        cal_rect = (ML, bot - ch, CW, ch)
        bot = bot - ch - 0.16
    else:
        cal_rect = None
    return top, bot, cal_rect


def callout(slide, spec, cal_rect, page):
    if not cal_rect:
        return
    cal = spec["callout"]
    x, y, w, h = cal_rect
    col = cal.get("color", "accent")
    fill_key = (col + "_l") if (col + "_l") in C else "accentl"
    rect(slide, x, y, w, h, fill=fill_key, line=None, rounded=True, radius=0.14)
    rect(slide, x, y, 0.055, h, fill=col, rounded=False)
    text(slide, x + 0.24, y, w - 0.46, h,
         [P(cal["text"], size=10.5, color="ink", line_spacing=1.18)],
         anchor="mid", page=page, where="callout")


def footnote(slide, spec, page, dark=False):
    fn = spec.get("footnote")
    if not fn:
        return
    text(slide, ML, FN_Y, CW, 0.3,
         [P("※ " + fn, size=9, color="8FA3B8" if dark else "muted",
            line_spacing=1.1)], page=page, where="footnote")


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# --------------------------------------------------------------- 深色装饰 ----
def deco_dark(slide):
    cx, cy = SW - 0.9, -0.55
    for d, col in ((5.6, "navy2"), (3.7, "1B3E63"), (2.0, "0F5F6B")):
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2),
                                     Inches(cy - d / 2), Inches(d), Inches(d))
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(col)
        shp.line.fill.background()
        shp.name = "deco-ring"          # 有意出血到画布外的装饰，自检时跳过
        _no_shadow(shp)


def bottom_bar(slide):
    y, h = SH - 0.13, 0.13
    for x, w, col in ((0, 5.4, "blue"), (5.4, 2.5, "accent"), (7.9, SW - 7.9, "green")):
        rect(slide, x, y, w, h, fill=col, rounded=False)


# ------------------------------------------------------------------- 封面 ----
def r_cover(prs, spec, page, total):
    s = new_slide(prs)
    bg(s, "navy")
    deco_dark(s)
    text(s, ML, 1.42, CW - 2.0, 0.32,
         [P(spec["kicker"], size=12, bold=True, color="accent", spc=150)],
         page=page, where="cover.kicker")
    text(s, ML, 1.92, CW - 1.2, 1.0,
         [P(spec["title"], size=44, bold=True, color="white", line_spacing=1.05)],
         page=page, where="cover.title")
    rect(s, ML, 3.02, 1.15, 0.055, fill="accent", rounded=False)
    text(s, ML, 3.28, CW - 1.6, 0.7,
         [P(spec["subtitle"], size=15, color="CBD6E2", line_spacing=1.25)],
         page=page, where="cover.subtitle")
    bw, gap, y = 2.05, 0.26, 4.42
    for i, b in enumerate(spec["badges"]):
        x = ML + i * (bw + gap)
        rect(s, x, y, bw, 0.54, fill="navy2", line="0F5F6B", lw=1.0, rounded=True,
             radius=0.5)
        text(s, x, y, bw, 0.54, [P(b, size=11.5, bold=True, color="white",
                                   align="center")], anchor="mid", page=page,
             where="cover.badge")
    text(s, ML, 6.30, CW, 0.3, [P(spec["meta"], size=10, color="8FA3B8")],
         page=page, where="cover.meta")
    bottom_bar(s)
    footer(s, page, total, dark=True)
    return s


# ------------------------------------------------------------------- 目录 ----
def r_toc(prs, spec, page, total):
    s = new_slide(prs)
    bg(s)
    top = header(s, spec, page)
    items = spec["items"]
    rh = min(0.94, (BODY_BOT - top) / len(items))
    for i, it in enumerate(items):
        y = top + i * rh
        rect(s, ML, y + 0.08, 0.66, rh - 0.26, fill="navy", rounded=True, radius=0.20)
        text(s, ML, y + 0.08, 0.66, rh - 0.26,
             [P(it["no"], size=17, bold=True, color="accent", align="center")],
             anchor="mid", page=page, where="toc.no")
        x2 = ML + 0.94
        text(s, x2, y + 0.14, 5.6, 0.32,
             [{"runs": [R(it["title"], size=14.5, bold=True, color="ink"),
                        R("   " + it["en"], size=10, color="accent")]}],
             page=page, where="toc.title")
        text(s, x2, y + 0.50, CW - 2.9, 0.34,
             [P(it["desc"], size=10.5, color="muted", line_spacing=1.15)],
             page=page, where="toc.desc")
        tw = 1.28
        rect(s, ML + CW - tw, y + 0.24, tw, 0.34, fill="accentl", rounded=True,
             radius=0.5)
        text(s, ML + CW - tw, y + 0.24, tw, 0.34,
             [P(it["n"], size=9.5, bold=True, color="accentd", align="center")],
             anchor="mid", page=page, where="toc.n")
        if i < len(items) - 1:
            line(s, x2, y + rh - 0.03, ML + CW, y + rh - 0.03, color="line")
    footer(s, page, total)
    return s


# ----------------------------------------------------------------- 章节页 ----
def r_divider(prs, spec, page, total):
    s = new_slide(prs)
    bg(s, "navy")
    deco_dark(s)
    text(s, ML - 0.06, 0.92, 3.4, 1.5,
         [P(spec["no"], size=88, bold=True, color="1D3E60", line_spacing=1.0)],
         page=page, where="div.no")
    text(s, ML, 2.46, CW - 1.4, 0.72,
         [P(spec["title"], size=34, bold=True, color="white", line_spacing=1.0)],
         page=page, where="div.title")
    text(s, ML, 3.24, CW - 1.4, 0.3,
         [P(spec["en"], size=12.5, bold=True, color="accent", spc=90)],
         page=page, where="div.en")
    rect(s, ML, 3.72, 0.9, 0.05, fill="accent", rounded=False)
    text(s, ML, 3.96, CW - 2.6, 0.7,
         [P(spec["lede"], size=12.5, color="CBD6E2", line_spacing=1.25)],
         page=page, where="div.lede")
    x, y = ML, 5.02
    for t in spec["terms"]:
        w = em_width(t) * 11 / 72 + 0.52
        if x + w > ML + CW:
            x, y = ML, y + 0.60
        rect(s, x, y, w, 0.44, fill="navy2", line="0F5F6B", lw=1.0, rounded=True,
             radius=0.5)
        text(s, x, y, w, 0.44, [P(t, size=11, bold=True, color="white",
                                  align="center")], anchor="mid", page=page,
             where="div.term")
        x += w + 0.24
    bottom_bar(s)
    footer(s, page, total, dark=True)
    return s


# --------------------------------------------------------------- 要点列表 ----
def r_points(prs, spec, page, total):
    s = new_slide(prs)
    bg(s)
    top = header(s, spec, page)
    _, bot, _ = body_geom(spec)
    items = spec["items"]
    rh = (bot - top) / len(items)
    for i, it in enumerate(items):
        y = top + i * rh
        rect(s, ML, y + 0.14, 0.52, rh - 0.36, fill="accentl", rounded=True,
             radius=0.22)
        text(s, ML, y + 0.14, 0.52, rh - 0.36,
             [P("%02d" % (i + 1), size=14, bold=True, color="accentd",
                align="center")], anchor="mid", page=page, where="points.no")
        text(s, ML + 0.80, y + 0.16, 2.5, 0.4,
             [P(it["head"], size=14.5, bold=True, color="ink")], page=page,
             where="points.head")
        text(s, ML + 3.48, y + 0.14, CW - 3.48, rh - 0.3,
             [P(it["body"], size=11.5, color="muted", line_spacing=1.3)],
             page=page, where="points.body")
        if i < len(items) - 1:
            line(s, ML + 0.80, y + rh - 0.05, ML + CW, y + rh - 0.05, color="line")
    footnote(s, spec, page)
    footer(s, page, total)
    return s


# ------------------------------------------------------------------- 卡片 ----
def _meter(slide, x, y, w, vals, page):
    """蓝绿权重条：vals = [蓝%, 绿%]"""
    h = 0.17
    b, g = vals
    rect(slide, x, y, w, h, fill="E7EDF4", rounded=True, radius=0.5)
    if b:
        rect(slide, x, y, max(w * b / 100.0, 0.12), h, fill="blue", rounded=True,
             radius=0.5)
    if g:
        gw = max(w * g / 100.0, 0.12)
        rect(slide, x + w - gw, y, gw, h, fill="green", rounded=True, radius=0.5)
    text(slide, x, y + h + 0.05, w, 0.26,
         [{"runs": [R("流量权重  ", size=8.5, color="muted"),
                    R("蓝 %d" % b, size=8.5, bold=True, color="blue"),
                    R("  :  ", size=8.5, color="muted"),
                    R("绿 %d" % g, size=8.5, bold=True, color="green")]}],
         page=page, where="card.meter")


def r_cards(prs, spec, page, total):
    s = new_slide(prs)
    bg(s)
    top = header(s, spec, page)
    _, bot, cal_rect = body_geom(spec)
    cards = spec["cards"]
    cols = spec.get("cols", len(cards))
    gap = 0.30
    cw = (CW - gap * (cols - 1)) / cols
    ch = bot - top
    fs = {2: {"title": 15, "en": 9.5, "body": 11.5, "tag": 9},
          3: {"title": 13.5, "en": 8.5, "body": 10.5, "tag": 8.5}}[cols]
    pad = 0.28 if cols == 2 else 0.24
    for i, cd in enumerate(cards):
        x = ML + i * (cw + gap)
        acc = cd.get("accent", "accent")
        rect(s, x, top, cw, ch, fill="card", line="line", lw=1.0, rounded=True,
             radius=0.045)
        rect(s, x + 0.10, top, cw - 0.20, 0.085, fill=acc, rounded=True, radius=0.5)
        tw = em_width(cd["tag"]) * fs["tag"] / 72 + 0.34 if cd.get("tag") else 0
        if cd.get("badge"):
            d = 0.46
            shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + cw - pad - d),
                                     Inches(top + 0.30), Inches(d), Inches(d))
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(acc)
            shp.line.fill.background()
            _no_shadow(shp)
            text(s, x + cw - pad - d, top + 0.30, d, d,
                 [P(cd["badge"], size=18, bold=True, color="white", align="center")],
                 anchor="mid", page=page, where="card.badge")
        elif cd.get("tag"):
            rect(s, x + cw - pad - tw, top + 0.30, tw, 0.32, fill=acc + "_l"
                 if (acc + "_l") in C else "accentl", rounded=True, radius=0.5)
            text(s, x + cw - pad - tw, top + 0.30, tw, 0.32,
                 [P(cd["tag"], size=fs["tag"], bold=True, color=acc, align="center")],
                 anchor="mid", page=page, where="card.tag")
        text(s, x + pad, top + 0.30, cw - 2 * pad - (tw or 0.6) - 0.1, 0.34,
             [P(cd["title"], size=fs["title"], bold=True, color="ink",
                line_spacing=1.05)], page=page, where="card.title")
        y = top + 0.30 + fs["title"] * 1.25 / 72 + 0.04
        text(s, x + pad, y, cw - 2 * pad, 0.24,
             [P(cd.get("en", ""), size=fs["en"], bold=True, color=acc, spc=30)],
             page=page, where="card.en")
        y += 0.30
        if cd.get("meter"):
            _meter(s, x + pad, y, cw - 2 * pad, cd["meter"], page)
            y += 0.52
        line(s, x + pad, y, x + cw - pad, y, color="line")
        y += 0.16
        paras = []
        for j, b in enumerate(cd["body"]):
            paras.append({"runs": [R("▪ ", size=fs["body"], bold=True, color=acc),
                                   R(b, size=fs["body"], color="ink")],
                          "line_spacing": 1.22,
                          "space_after": 0 if j == len(cd["body"]) - 1 else 7})
        text(s, x + pad, y, cw - 2 * pad, top + ch - y - 0.20, paras, page=page,
             where="card.body")
    callout(s, spec, cal_rect, page)
    footnote(s, spec, page)
    footer(s, page, total)
    return s


# ----------------------------------------------------------------- 时间轴 ----
def r_timeline(prs, spec, page, total):
    s = new_slide(prs)
    bg(s)
    top = header(s, spec, page)
    _, bot, cal_rect = body_geom(spec)
    steps = spec["steps"]
    n = len(steps)
    big = n <= 3
    if spec.get("horizontal"):
        gap = 0.34
        cw = (CW - gap * (n - 1)) / n
        line(s, ML + cw / 2, top + 0.33, ML + CW - cw / 2, top + 0.33, color="line",
             w=1.5)
        for i, st in enumerate(steps):
            x = ML + i * (cw + gap)
            d = 0.66
            shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + cw / 2 - d / 2),
                                     Inches(top), Inches(d), Inches(d))
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb("navy")
            shp.line.color.rgb = rgb("accent")
            shp.line.width = Pt(1.5)
            _no_shadow(shp)
            text(s, x + cw / 2 - d / 2, top, d, d,
                 [P(st["tag"], size=17, bold=True, color="white", align="center")],
                 anchor="mid", page=page, where="tl.tag")
            text(s, x, top + 0.88, cw, 0.34,
                 [P(st["title"], size=14, bold=True, color="ink", align="center")],
                 page=page, where="tl.title")
            text(s, x + 0.06, top + 1.28, cw - 0.12, bot - top - 1.3,
                 [P(st["body"], size=11.5, color="muted", align="center",
                    line_spacing=1.3)], page=page, where="tl.body")
    else:
        rh = (bot - top) / n
        ts = 13.5 if big else 12.0
        bs = 11.5 if big else 10.5
        line(s, ML + 0.37, top + 0.24, ML + 0.37, top + rh * (n - 1) + 0.24,
             color="line", w=1.5)
        for i, st in enumerate(steps):
            y = top + i * rh
            rect(s, ML, y + 0.07, 0.74, 0.34, fill="navy", rounded=True, radius=0.5)
            text(s, ML, y + 0.07, 0.74, 0.34,
                 [P(st["tag"], size=10.5, bold=True, color="accent", align="center")],
                 anchor="mid", page=page, where="tl.tag")
            text(s, ML + 1.02, y + 0.07, 2.72, 0.34,
                 [P(st["title"], size=ts, bold=True, color="ink")], page=page,
                 where="tl.title")
            text(s, ML + 3.90, y + 0.05, CW - 3.90, rh - 0.14,
                 [P(st["body"], size=bs, color="muted", line_spacing=1.26)],
                 page=page, where="tl.body")
            if i < n - 1:
                line(s, ML + 1.02, y + rh - 0.06, ML + CW, y + rh - 0.06, color="line")
    callout(s, spec, cal_rect, page)
    footnote(s, spec, page)
    footer(s, page, total)
    return s


# ------------------------------------------------------------------- 表格 ----
def _neutral_table_style(tbl):
    """切到 "No Style, Table Grid"，让显式设置的填充/边框说了算。"""
    tblPr = tbl._tbl.find(qn("a:tblPr"))
    if tblPr is None:
        return
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(sid)
    sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"


def _cell_border(cell, color="line", w_pt=0.75, edges=("L", "R", "T", "B")):
    tcPr = cell._tc.get_or_add_tcPr()
    for e in reversed(edges):
        tag = qn("a:ln" + e)
        old = tcPr.find(tag)
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement(tag, {"w": str(int(w_pt * 12700)), "cap": "flat",
                                    "cmpd": "sng", "algn": "ctr"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = fill.makeelement(qn("a:srgbClr"), {"val": C.get(color, color)})
        fill.append(clr)
        ln.append(fill)
        tcPr.insert(0, ln)


def _cell_text(cell, runs, size=10.5, bold=False, color="ink", align="left",
               anchor="mid", page=0, where=""):
    tf = cell.text_frame
    tf.word_wrap = True
    cell.margin_left = cell.margin_right = Inches(0.10)
    cell.margin_top = cell.margin_bottom = Inches(0.045)
    cell.vertical_anchor = ANCHOR[anchor]
    par = tf.paragraphs[0]
    par.alignment = ALIGN[align]
    par.line_spacing = 1.18
    if not runs:
        runs = [{"t": ""}]
    for r in runs:
        run = par.add_run()
        run.text = r.get("t", "")
        f = run.font
        f.name = FONT
        f.size = Pt(r.get("size", size))
        f.bold = r.get("bold", bold)
        f.color.rgb = rgb(r.get("color", color))
        _set_ea(run, FONT)
    return par


def r_table(prs, spec, page, total):
    s = new_slide(prs)
    bg(s)
    top = header(s, spec, page)
    heads, rows, widths = spec["headers"], spec["rows"], spec["widths"]
    compact = spec.get("compact", False)
    tag_col = spec.get("tag_col")
    row_colors = spec.get("row_colors")
    scale = CW / float(sum(widths))
    n = len(rows)
    avail = BODY_BOT - top
    head_h = 0.40
    row_h = min(0.60 if not compact else 0.60, (avail - head_h) / n)
    th = head_h + row_h * n
    ty = top + max(0.0, (avail - th) / 2)
    hs = 10 if compact else 11
    bs = 9.5 if compact else 11
    gf = s.shapes.add_table(n + 1, len(heads), Inches(ML), Inches(ty), Inches(CW),
                            Inches(th))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    _neutral_table_style(tbl)
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w * scale)
    tbl.rows[0].height = Inches(head_h)
    for i in range(n):
        tbl.rows[i + 1].height = Inches(row_h)

    # 单元格文字高度自检：PowerPoint 会按内容自动撑高行，这里提前发现
    for j, wj in enumerate(widths):
        colw = wj * scale - 0.20
        for i, row in enumerate([heads] + rows):
            val = DOMAINS[row[j]]["label"] if (tag_col is not None and j == tag_col
                                               and i > 0) else str(row[j])
            sz = hs if i == 0 else (bs - 0.5 if (tag_col is not None and
                                                  j == tag_col) else bs)
            need = measure([{"runs": [{"t": val, "size": sz}]}], colw) + 0.09
            have = head_h if i == 0 else row_h
            FIT_LOG.append({"page": page, "where": "table.r%dc%d" % (i, j),
                            "w": round(colw, 2), "h": round(have, 2),
                            "need": round(need, 2), "over": round(need - have, 2),
                            "head": val[:30]})

    for j, htxt in enumerate(heads):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb("navy")
        _cell_border(cell, color="navy", w_pt=0.75)
        _cell_text(cell, [{"t": htxt}], size=hs, bold=True, color="white",
                   align="center" if (compact and j == len(heads) - 1) else "left")
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("card" if i % 2 == 0 else "F2F6FA")
            _cell_border(cell, color="line", w_pt=0.75)
            if tag_col is not None and j == tag_col:
                dom = DOMAINS.get(val, {"label": val, "color": "muted"})
                _cell_text(cell, [{"t": dom["label"]}], size=bs - 0.5, bold=True,
                           color=dom["color"], align="center")
            elif j == 0:
                col = "ink"
                if row_colors and row_colors[i]:
                    col = row_colors[i]
                _cell_text(cell, [{"t": val}], size=bs, bold=True, color=col)
            else:
                _cell_text(cell, [{"t": val}], size=bs,
                           color="muted" if not compact else "ink",
                           bold=False)
    footnote(s, spec, page)
    footer(s, page, total)
    return s


# --------------------------------------------------------------- 拓扑图 ----
_anchor_pts = anchor_points        # 复用 deck_content 的共用实现


def r_diagram(prs, spec, page, total):
    s = new_slide(prs)
    bg(s)
    top = header(s, spec, page)
    _, bot, cal_rect = body_geom(spec)
    d = spec["diagram"]
    avail_h = bot - top
    w = min(CW, avail_h * DW / DH)
    h = w * DH / DW
    x0 = (SW - w) / 2.0
    y0 = top + max(0.0, (avail_h - h) / 2.0)
    k = w / float(DW)
    X = lambda v: x0 + v * k
    Y = lambda v: y0 + v * k
    nodes = {n["id"]: n for n in d["nodes"]}

    # 连线（先画，压在节点下面）
    for e in d["edges"]:
        (sx, sy), (tx, ty) = _anchor_pts(nodes[e["from"]], nodes[e["to"]])
        dash = MSO_LINE_DASH_STYLE.DASH if e.get("kind") == "dashed" else None
        arrow(s, X(sx), Y(sy), X(tx), Y(ty), color=e.get("color", "navy2"),
              w=1.25, dash=dash)
        if e.get("label"):
            mx, my = (sx + tx) / 2.0, (sy + ty) / 2.0
            lw = em_width(e["label"]) * 8.5 / 72 + 0.16
            rect(s, X(mx) - lw / 2, Y(my) - 0.09, lw, 0.18, fill="bg", rounded=True,
                 radius=0.5)
            text(s, X(mx) - lw / 2, Y(my) - 0.09, lw, 0.18,
                 [P(e["label"], size=8.5, bold=True,
                    color=e.get("color", "navy2"), align="center")], anchor="mid",
                 page=page, where="dia.edge")

    # 图例 + 右上角说明
    ly = 2
    for i, item in enumerate(d.get("legend", [])):
        yy = ly + i * 20
        dash = MSO_LINE_DASH_STYLE.DASH if item["kind"] == "dashed" else None
        col = "muted" if item["kind"] == "dashed" else "navy2"
        line(s, X(20), Y(yy + 8), X(54), Y(yy + 8), color=col, w=1.25, dash=dash)
        text(s, X(62), Y(yy), 2.8, 0.22, [P(item["text"], size=8.5, color="muted")],
             page=page, where="dia.legend")
    cap = "两站同机房 · 同一内网拓扑 · RTT 对等（延迟忽略不计）"
    text(s, X(770), Y(ly), X(1180) - X(770), 0.24,
         [P(cap, size=8.5, bold=True, color="accentd", align="right")],
         page=page, where="dia.cap")

    # 节点
    for n in d["nodes"]:
        st = n["style"]
        zone = n.get("zone")
        x, y, nw, nh = X(n["x"]), Y(n["y"]), n["w"] * k, n["h"] * k
        if st == "band":
            text(s, x, y, nw, nh * 1.2,
                 [P(n["title"], size=10, bold=True, color="muted", align="center")],
                 anchor="mid", page=page, where="dia.band")
            continue
        if st == "zone":
            rect(s, x, y, nw, nh, fill=zone + "_l", line=zone, lw=1.25, rounded=True,
                 radius=0.05)
            tcol, bcol = zone, "ink"
            tsz, bsz = 11.5, 9
        elif st == "actor":
            rect(s, x, y, nw, nh, fill="navy", line=None, rounded=True, radius=0.30)
            tcol, bcol, tsz, bsz = "white", "CBD6E2", 10.5, 8
        elif st == "gateway":
            rect(s, x, y, nw, nh, fill="navy2", line="accent", lw=1.5, rounded=True,
                 radius=0.16)
            tcol, bcol, tsz, bsz = "white", "9FE3EC", 11.5, 9
        elif st == "inner":
            rect(s, x, y, nw, nh, fill="card", line=zone, lw=1.0, rounded=True,
                 radius=0.10)
            tcol, bcol, tsz, bsz = "ink", "muted", 10, 8
        elif st == "registry":
            rect(s, x, y, nw, nh, fill="card", line="B7C4D2", lw=1.0, rounded=True,
                 radius=0.12, dash=MSO_LINE_DASH_STYLE.DASH)
            tcol, bcol, tsz, bsz = "accentd", "muted", 9.5, 7.5
        else:                                    # shared
            rect(s, x, y, nw, nh, fill="EDF2F8", line="B9C9D9", lw=1.0, rounded=True,
                 radius=0.07)
            tcol, bcol, tsz, bsz = "ink", "muted", 10.5, 8.5
        pad = 0.11 if st in ("zone", "shared") else 0.09
        align = "center" if st in ("actor", "gateway", "band") else "left"
        paras = [P(n["title"], size=tsz, bold=True, color=tcol, align=align,
                   line_spacing=1.06)]
        if n.get("body"):
            paras.append(P(n["body"], size=bsz, color=bcol, align=align,
                           line_spacing=1.12, space_before=2))
        text(s, x + pad, y + pad * 0.7, nw - 2 * pad, nh - pad, paras,
             anchor="top" if st == "zone" else "mid", page=page,
             where="dia." + n["id"])
    callout(s, spec, cal_rect, page)
    footnote(s, spec, page)
    footer(s, page, total)
    return s


# --------------------------------------------------------------- 检查清单 ----
def r_checklist(prs, spec, page, total):
    s = new_slide(prs)
    bg(s)
    top = header(s, spec, page)
    _, bot, cal_rect = body_geom(spec)
    groups = spec["groups"]
    gap = 0.30
    cw = (CW - gap * (len(groups) - 1)) / len(groups)
    ch = bot - top
    for i, g in enumerate(groups):
        x = ML + i * (cw + gap)
        acc = g.get("accent", "accent")
        rect(s, x, top, cw, ch, fill="card", line="line", lw=1.0, rounded=True,
             radius=0.04)
        rect(s, x, top, cw, 0.46, fill=acc, rounded=True, radius=0.10)
        text(s, x + 0.22, top, cw - 0.4, 0.46,
             [{"runs": [R(g["title"], size=12.5, bold=True, color="white"),
                        R("   %d 项" % len(g["items"]), size=9.5, color="FFFFFF")]}],
             anchor="mid", page=page, where="ck.head")
        paras = []
        for j, it in enumerate(g["items"]):
            paras.append({"runs": [R("□ ", size=11, bold=True, color=acc),
                                   R(it, size=10.5, color="ink")],
                          "line_spacing": 1.2,
                          "space_after": 0 if j == len(g["items"]) - 1 else 8})
        text(s, x + 0.22, top + 0.66, cw - 0.42, ch - 0.86, paras, page=page,
             where="ck.items")
    callout(s, spec, cal_rect, page)
    footnote(s, spec, page)
    footer(s, page, total)
    return s


# ----------------------------------------------------------------- 结尾页 ----
def r_closing(prs, spec, page, total):
    s = new_slide(prs)
    bg(s, "navy")
    deco_dark(s)
    text(s, ML, 1.62, CW - 1.5, 0.8,
         [P(spec["title"], size=36, bold=True, color="white", line_spacing=1.0)],
         page=page, where="end.title")
    rect(s, ML, 2.62, 1.0, 0.055, fill="accent", rounded=False)
    text(s, ML, 2.86, CW - 2.0, 0.4,
         [P(spec["subtitle"], size=13, color="CBD6E2")], page=page,
         where="end.subtitle")
    n = len(spec["stats"])
    gap = 0.34
    cw = (CW - 2.2 - gap * (n - 1)) / n
    y, chh = 3.86, 1.72
    for i, st in enumerate(spec["stats"]):
        x = ML + i * (cw + gap)
        rect(s, x, y, cw, chh, fill="navy2", line="0F5F6B", lw=1.0, rounded=True,
             radius=0.07)
        rect(s, x + 0.24, y + 0.28, 0.42, 0.05, fill="accent", rounded=False)
        text(s, x + 0.24, y + 0.44, cw - 0.48, 0.6,
             [P(st["num"], size=27, bold=True, color="white", line_spacing=1.0)],
             page=page, where="end.num")
        text(s, x + 0.24, y + 1.06, cw - 0.48, 0.55,
             [P(st["label"], size=10.5, color="9FB4C9", line_spacing=1.25)],
             page=page, where="end.label")
    text(s, ML, 6.24, CW, 0.3, [P(spec["meta"], size=10, color="8FA3B8")],
         page=page, where="end.meta")
    bottom_bar(s)
    footer(s, page, total, dark=True)
    return s


# --------------------------------------------------------- 演讲者备注生成 ----
def add_notes(slide, spec):
    bits = []
    for key in ("kicker", "title", "lede"):
        if spec.get(key):
            bits.append(spec[key])
    for it in spec.get("items", []):
        if isinstance(it, dict) and "head" in it:
            bits.append("%s：%s" % (it["head"], it["body"]))
        elif isinstance(it, dict) and "title" in it:
            bits.append("%s（%s）：%s" % (it["title"], it.get("en", ""),
                                          it.get("desc", "")))
    for cd in spec.get("cards", []):
        bits.append("%s（%s）：%s" % (cd["title"], cd.get("en", ""),
                                      "；".join(cd.get("body", []))))
    for st in spec.get("steps", []):
        bits.append("%s %s：%s" % (st["tag"], st["title"], st["body"]))
    for row in spec.get("rows", []):
        bits.append(" | ".join(str(c) for c in row))
    for g in spec.get("groups", []):
        bits.append("%s：%s" % (g["title"], "；".join(g["items"])))
    for st in spec.get("stats", []):
        bits.append("%s = %s" % (st["num"], st["label"].replace("\n", " ")))
    if spec.get("callout"):
        bits.append("【强调】" + spec["callout"]["text"])
    if spec.get("footnote"):
        bits.append("【备注】" + spec["footnote"])
    bits.append("口径来源：CONTEXT.md（体检系统升级V2 领域术语表）。")
    tf = slide.notes_slide.notes_text_frame
    tf.text = "\n".join(bits)


RENDERERS = {
    "cover": r_cover, "toc": r_toc, "divider": r_divider, "points": r_points,
    "cards": r_cards, "timeline": r_timeline, "table": r_table,
    "diagram": r_diagram, "checklist": r_checklist, "closing": r_closing,
}


def build(out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    total = len(SLIDES)
    for i, spec in enumerate(SLIDES, 1):
        kind = spec["kind"]
        if kind not in RENDERERS:
            raise ValueError("未实现的版式：%s" % kind)
        slide = RENDERERS[kind](prs, spec, i, total)
        add_notes(slide, spec)
    cp = prs.core_properties
    cp.title = META["title"]
    cp.author = META["author"]
    cp.subject = META["subject"]
    cp.comments = META["comments"]
    cp.category = "架构评审 / 领域术语"
    cp.keywords = "蓝绿双站;原子切流;Expand and Contract;区域亲和路由;体检系统升级V2"
    prs.save(out_path)
    return out_path
